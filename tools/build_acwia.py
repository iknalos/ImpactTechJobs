"""Build data/acwia_employers.json from DOL prevailing-wage disclosure files.

Source: DOL OFLC "Prevailing Wage Program, Form ETA-9141" public disclosure
xlsx (https://www.dol.gov/agencies/eta/foreign-labor/performance). Each PWD
request answers "Is the employer covered by ACWIA?" with sub-boxes for
institution of higher education / affiliated nonprofit / nonprofit or
governmental research organization: the same four categories that make an
employer H-1B cap-exempt (INA 214(g)(5)(A)-(B), 8 CFR 214.2(h)(8)(iii)(F)).
PWD_WAGE_SOURCE = "OES (ACWIA)" means DOL itself issued the higher-ed wage,
i.e. accepted the claim.

dol.gov blocks scripted downloads, so the files are fetched by hand in a
browser (a few times a year is plenty) and passed in here:

    python tools/build_acwia.py ~/Downloads/PW_Disclosure_Data_FY2025_Q4.xlsx ...

Rows are streamed (openpyxl read-only); only ACWIA-covered employers are kept.
"""

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

OUT = Path(__file__).resolve().parent.parent / "data" / "acwia_employers.json"

WANT = {
    "name": ("EMPLOYER_LEGAL_BUSINESS_NAME", "EMPLOYER_NAME"),
    "dba": ("TRADE_NAME_DBA",),
    "city": ("EMPLOYER_CITY",),
    "state": ("EMPLOYER_STATE",),
    "fein": ("EMPLOYER_FEIN",),
    "naics": ("NAICS_CODE",),
    "acwia": ("COVERED_BY_ACWIA",),
    "hied": ("ACWIA_INST_HIGHER_EDUCATION",),
    "affil": ("ACWIA_AFFILIATED_NON_PROFIT",),
    "research": ("ACWIA_RESEARCH_ORG",),
    "source": ("PWD_WAGE_SOURCE",),
    "date": ("DETERMINATION_DATE", "PREVAIL_WAGE_DETERM_DATE", "RECEIVED_DATE"),
}


def yes(v):
    return str(v or "").strip().upper() in ("Y", "YES", "TRUE")


def fein_key(v):
    d = re.sub(r"\D", "", str(v or ""))
    return d if len(d) == 9 else None


def read(path, emps):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [str(h or "").strip().upper() for h in next(rows)]
    col = {}
    for k, names in WANT.items():
        col[k] = next((header.index(n) for n in names if n in header), None)
    missing = [k for k in ("name", "acwia") if col[k] is None]
    if missing:
        raise SystemExit("%s: missing columns %s" % (path, missing))
    fy = re.search(r"FY(\d{4})", Path(path).name)
    fy = fy.group(1) if fy else "?"
    n = kept = 0
    for r in rows:
        n += 1
        g = lambda k: r[col[k]] if col[k] is not None and col[k] < len(r) else None  # noqa: E731
        if not yes(g("acwia")):
            continue
        name = str(g("name") or "").strip()
        if not name:
            continue
        kept += 1
        key = fein_key(g("fein")) or "name:" + name.upper()
        e = emps.setdefault(key, {
            "fein": fein_key(g("fein")), "names": {}, "dbas": {}, "city": None, "state": None,
            "naics": None, "filings": 0, "higher_ed": 0, "affiliated_nonprofit": 0,
            "research_org": 0, "dol_acwia_wage": 0, "fy": set(), "last": None})
        e["names"][name] = e["names"].get(name, 0) + 1
        dba = str(g("dba") or "").strip()
        if dba and dba.upper() not in ("N/A", "NA", "NONE"):
            e["dbas"][dba] = e["dbas"].get(dba, 0) + 1
        e["city"] = e["city"] or (str(g("city") or "").strip() or None)
        e["state"] = e["state"] or (str(g("state") or "").strip() or None)
        e["naics"] = e["naics"] or (str(g("naics") or "").strip() or None)
        e["filings"] += 1
        e["higher_ed"] += yes(g("hied"))
        e["affiliated_nonprofit"] += yes(g("affil"))
        e["research_org"] += yes(g("research"))
        e["dol_acwia_wage"] += "ACWIA" in str(g("source") or "").upper()
        e["fy"].add(fy)
        d = g("date")
        d = d.date().isoformat() if hasattr(d, "date") else (str(d)[:10] if d else None)
        if d and (not e["last"] or d > e["last"]):
            e["last"] = d
    wb.close()
    print("%s: %d rows, %d ACWIA-covered" % (Path(path).name, n, kept))


def main(paths):
    if not paths:
        raise SystemExit(__doc__)
    emps = {}
    for p in paths:
        read(p, emps)
    out = []
    for e in emps.values():
        names = sorted(e["names"], key=e["names"].get, reverse=True)
        cats = [c for c in ("higher_ed", "affiliated_nonprofit", "research_org") if e[c]]
        out.append({
            "name": names[0], "aliases": names[1:6],
            "dba": sorted(e["dbas"], key=e["dbas"].get, reverse=True)[:3],
            "fein_last4": e["fein"][-4:] if e["fein"] else None,
            "city": e["city"], "state": e["state"], "naics": e["naics"],
            "filings": e["filings"], "dol_acwia_wage": e["dol_acwia_wage"],
            "basis": max(cats, key=lambda c: e[c]) if cats else None,
            "fy": sorted(e["fy"]), "last": e["last"],
        })
    out.sort(key=lambda x: -x["filings"])
    OUT.parent.mkdir(exist_ok=True)
    OUT.write_text(json.dumps({"source": [Path(p).name for p in paths], "employers": out},
                              separators=(",", ":")), encoding="utf-8")
    print("Wrote %d ACWIA-covered employers -> %s" % (len(out), OUT))


if __name__ == "__main__":
    main(sys.argv[1:])
